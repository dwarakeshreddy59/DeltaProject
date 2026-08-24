"""
export_excel.py – Generate a styled Excel workbook from extracted records.
"""

import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Column definitions ─────────────────────────────────────────────────────

PO_COLS = [
    ("po_number",       "PO Number"),
    ("po_date",         "PO Date"),
    ("po_description",  "PO Description"),
    ("delivery_date",   "Delivery Date / PO Validity"),
    ("total_amount",    "Total PO Amount"),
]

INVOICE_COLS = [
    ("invoice_number",      "Invoice Number"),
    ("invoice_date",        "Invoice Date"),
    ("po_number",           "PO Number (Ref)"),
    ("inv_description",     "Invoice Description"),
    ("invoice_period",      "Invoice Period"),
    ("assessable_value",    "Assessable Value"),
    ("gst_rate",            "GST Rate (%)"),
    ("gst_amount",          "GST Amount"),
    ("total_tax",           "Total Tax"),
    ("total_invoice_value", "Total Invoice Value"),
    ("tds_rate",            "TDS Rate (%)"),
    ("tds_amount",          "TDS Amount"),
    ("receivable",          "Receivable"),
]

REMITTANCE_COLS = [
    ("remittance_number",   "Remittance Number"),
    ("remittance_date",     "Remittance Date"),
    ("invoice_number",      "Invoice Number (Ref)"),
    ("gross_amount",        "Gross Amount"),
    ("total_gross_amount",  "Total Gross Amount"),
]


# ── Styles ─────────────────────────────────────────────────────────────────

HEADER_FONT   = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")   # dark blue
SUBHDR_FILL   = PatternFill("solid", fgColor="2E75B6")   # medium blue
CENTER        = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)
NUMBER_FMT    = "#,##0.00"


def _write_sheet(ws, title: str, rows: list, col_defs: list, fill: PatternFill) -> None:
    """Write a sheet with a title row, header row, and data rows."""
    keys   = [c[0] for c in col_defs]
    labels = [c[1] for c in col_defs]
    ncols  = len(labels)

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font      = Font(bold=True, color="FFFFFF", size=13)
    title_cell.fill      = fill
    title_cell.alignment = CENTER

    # Header row
    for col_idx, label in enumerate(labels, start=1):
        cell = ws.cell(row=2, column=col_idx, value=label)
        cell.font      = HEADER_FONT
        cell.fill      = SUBHDR_FILL
        cell.alignment = CENTER
        cell.border    = THIN_BORDER

    # Data rows
    for row_idx, record in enumerate(rows, start=3):
        for col_idx, key in enumerate(keys, start=1):
            val  = record.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if isinstance(val, (int, float)) and val not in ("", None):
                cell.number_format = NUMBER_FMT

    # Auto column width
    for col_idx in range(1, ncols + 1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, len(rows) + 3)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 25


def generate(records: list) -> str:
    """
    Build a styled Excel workbook with three sheets.
    Returns the path to the temp file.
    """
    wb = Workbook()

    # Sheet 1 – PO
    ws_po = wb.active
    ws_po.title = "Purchase Orders"
    _write_sheet(
        ws_po, "Purchase Order Details", records, PO_COLS,
        PatternFill("solid", fgColor="1F4E79"),
    )

    # Sheet 2 – Invoices
    ws_inv = wb.create_sheet("Invoices")
    _write_sheet(
        ws_inv, "Invoice Details & Calculations", records, INVOICE_COLS,
        PatternFill("solid", fgColor="375623"),
    )

    # Sheet 3 – Remittances
    ws_rem = wb.create_sheet("Remittances")
    _write_sheet(
        ws_rem, "Remittance Details", records, REMITTANCE_COLS,
        PatternFill("solid", fgColor="7B3F00"),
    )

    # Save to temp file
    tmp = tempfile.NamedTemporaryFile(
        suffix=".xlsx", delete=False,
        dir=tempfile.gettempdir(), prefix="pdf_portal_",
    )
    wb.save(tmp.name)
    tmp.close()
    return tmp.name
